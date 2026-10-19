import os
import re
import math
import operator
import ast
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import pandas as pd
import openpyxl

from Copytoalg import (
    parse_act_txt_folder,
    populate_simple_thresholds_template,
    update_alg_file,
    round_threshold_value
)

# ============================================================
# SAFE FORMULA EVALUATOR
# ============================================================
def normalize_formula(formula: str) -> str:
    formula = formula.strip().replace("^", "**")
    formula = re.sub(r"√\s*\(([^()]*)\)", r"sqrt(\1)", formula)
    return formula.replace("math.sqrt", "sqrt").replace("np.sqrt", "sqrt")

class SafeFormulaEvaluator:
    binary_operators = {
        ast.Add: operator.add,
        ast.Sub: operator.sub,
        ast.Mult: operator.mul,
        ast.Div: operator.truediv,
        ast.FloorDiv: operator.floordiv,
        ast.Mod: operator.mod,
        ast.Pow: operator.pow
    }
    unary_operators = {
        ast.UAdd: operator.pos,
        ast.USub: operator.neg
    }
    allowed_functions = {
        "sqrt": math.sqrt,
        "abs": abs,
        "min": min,
        "max": max,
        "round": round,
        "sin": math.sin,
        "cos": math.cos,
        "tan": math.tan,
        "log": math.log,
        "exp": math.exp
    }

    def __init__(self, variables: dict):
        self.variables = variables

    def evaluate(self, formula: str):
        parsed = ast.parse(formula, mode="eval")
        return self._evaluate_node(parsed.body)

    def _evaluate_node(self, node):
        if isinstance(node, ast.Constant):
            if isinstance(node.value, (int, float)):
                return float(node.value)
            raise ValueError("Numeric values only.")
        if isinstance(node, ast.Name):
            if node.id in self.variables:
                return float(self.variables[node.id])
            raise ValueError(f"Variable '{node.id}' not allowed. Use A, B, or C.")
        if isinstance(node, ast.BinOp):
            op = type(node.op)
            if op not in self.binary_operators:
                raise ValueError(f"Operator {op.__name__} not supported.")
            return self.binary_operators[op](self._evaluate_node(node.left), self._evaluate_node(node.right))
        if isinstance(node, ast.UnaryOp):
            op = type(node.op)
            if op not in self.unary_operators:
                raise ValueError(f"Unary operator {op.__name__} not supported.")
            return self.unary_operators[op](self._evaluate_node(node.operand))
        if isinstance(node, ast.Call):
            if not isinstance(node.func, ast.Name) or node.func.id not in self.allowed_functions:
                raise ValueError("Function not supported.")
            args = [self._evaluate_node(a) for a in node.args]
            return self.allowed_functions[node.func.id](*args)
        raise ValueError(f"Unsupported syntax: {type(node).__name__}")

def calculate_formula(formula: str, A: float, B: float, C: float) -> float:
    norm = normalize_formula(formula)
    evaluator = SafeFormulaEvaluator({"A": A, "B": B, "C": C})
    return evaluator.evaluate(norm)



# ============================================================
# PARAMETER NAME NORMALIZATION HELPER
# ============================================================
def clean_parameter_name(raw: str) -> str:
    """Strips leading number prefixes and condition tags."""
    if not raw or pd.isna(raw):
        return ""
    s = re.sub(r"^\d+_", "", str(raw)).strip()
    return re.sub(r"_(OFF|ON_LEFT|ON_RIGHT|ON)$", "", s, flags=re.IGNORECASE).strip()

# ============================================================
# GUI ROW COMPONENT FOR CHANNELS
# ============================================================
class ChannelRow:
    def __init__(self, parent, channel, max_val, min_val, row_idx, global_mode="OFF"):
        self.channel = str(channel)
        self.max_value = float(max_val)
        self.min_value = float(min_val)

        self.status_var = tk.StringVar(value=global_mode)
        init_c = self.max_value if global_mode == "OFF" else self.min_value
        self.c_var = tk.StringVar(value=str(init_c))
        self.a_var = tk.StringVar(value="10")
        self.b_var = tk.StringVar(value="0")
        self.formula1_var = tk.StringVar(value="sqrt(A^2+C^2)")
        self.formula2_var = tk.StringVar(value="")
        self.formula3_var = tk.StringVar(value="")
        self.result_var = tk.StringVar(value="")

        # Widgets
        pad = 2
        ttk.Label(parent, text=self.channel, font=("Segoe UI", 8, "bold")).grid(row=row_idx, column=0, padx=4, pady=pad, sticky="w")
        ttk.Label(parent, text=str(self.max_value)).grid(row=row_idx, column=1, padx=3, pady=pad)
        ttk.Label(parent, text=str(self.min_value)).grid(row=row_idx, column=2, padx=3, pady=pad)

        status_combo = ttk.Combobox(parent, textvariable=self.status_var, values=["OFF", "ON"], width=5, state="readonly")
        status_combo.grid(row=row_idx, column=3, padx=3, pady=pad)
        status_combo.bind("<<ComboboxSelected>>", self.on_status_change)

        ttk.Entry(parent, textvariable=self.c_var, width=8, state="readonly").grid(row=row_idx, column=4, padx=3, pady=pad)
        ttk.Entry(parent, textvariable=self.a_var, width=6).grid(row=row_idx, column=5, padx=3, pady=pad)
        ttk.Entry(parent, textvariable=self.b_var, width=6).grid(row=row_idx, column=6, padx=3, pady=pad)
        ttk.Entry(parent, textvariable=self.formula1_var, width=16).grid(row=row_idx, column=7, padx=3, pady=pad)
        ttk.Entry(parent, textvariable=self.formula2_var, width=16).grid(row=row_idx, column=8, padx=3, pady=pad)
        ttk.Entry(parent, textvariable=self.formula3_var, width=16).grid(row=row_idx, column=9, padx=3, pady=pad)
        ttk.Entry(parent, textvariable=self.result_var, width=10, state="readonly").grid(row=row_idx, column=10, padx=3, pady=pad)
        ttk.Button(parent, text="Calc", width=5, command=self.calculate).grid(row=row_idx, column=11, padx=3, pady=pad)

    def on_status_change(self, event=None):
        mode = self.status_var.get().strip().upper()
        if mode == "OFF":
            self.c_var.set(str(self.max_value))
        else:
            self.c_var.set(str(self.min_value))
        self.result_var.set("")

    def set_mode(self, mode):
        self.status_var.set(mode)
        self.on_status_change()

    def calculate(self, show_error=True):
        try:
            a_val = self.a_var.get().strip()
            b_val = self.b_var.get().strip()
            c_val = self.c_var.get().strip()

            if not a_val or not b_val or not c_val:
                raise ValueError("Values for A, B, and C are required.")

            A = float(a_val)
            B = float(b_val)
            C = float(c_val)

            formulas = [
                self.formula1_var.get().strip(),
                self.formula2_var.get().strip(),
                self.formula3_var.get().strip()
            ]
            active_formulas = [f for f in formulas if f]

            if not active_formulas:
                raise ValueError("At least one formula (Formula 1, 2, or 3) must be provided.")

            calculated_results = []
            for f in active_formulas:
                raw_eval = calculate_formula(f, A, B, C)
                # Apply round up to next integer if fractional part > 0.1
                rounded_val = round_threshold_value(raw_eval)
                calculated_results.append((f, raw_eval, rounded_val))

            # Pick highest rounded result among the provided formulas
            best_calc = max(calculated_results, key=lambda item: item[2])
            final_rounded_res = best_calc[2]

            disp_str = str(int(final_rounded_res)) if float(final_rounded_res).is_integer() else f"{final_rounded_res:.4g}"
            self.result_var.set(disp_str)

            return {
                "STATUS": "OK",
                "CHANNEL": self.channel,
                "C": C,
                "A": A,
                "B": B,
                "FORMULA": best_calc[0],
                "FORMULA1": formulas[0],
                "FORMULA2": formulas[1],
                "FORMULA3": formulas[2],
                "BEST_FORMULA": best_calc[0],
                "RAW_RESULT": best_calc[1],
                "RESULT": final_rounded_res
            }
        except Exception as e:
            self.result_var.set("ERR")
            if show_error:
                messagebox.showerror("Calculation Error", f"Channel {self.channel}: {e}")
            return {"STATUS": "ERROR", "CHANNEL": self.channel, "ERROR": str(e)}

# ============================================================
# MAIN APPLICATION
# ============================================================
class SignalMinMaxTemplateToolApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Integrated ACT Threshold & Signal MinMax Tool")
        self.root.geometry("1000x550")

        # Paths & Configurations
        self.act_folder_path = tk.StringVar(value=os.path.join(".", "OUTPUT", "ACT"))
        self.signal_file_path = tk.StringVar(value="outputchannelmaxmin.xlsx")
        self.template_path = tk.StringVar(value="act_result_output_template.xlsx")
        self.alg_path = tk.StringVar(value="input_parameter.alg")
        self.cal_mode = tk.StringVar(value="OFF")
        self.status_msg = tk.StringVar(value="Ready. Select input files/folders.")

        self.channel_rows = []
        self.decided_map = {}
        self.raw_val_map = {}

        self.create_widgets()

    def create_widgets(self):
        # Top Panel: Files & Configuration
        top_frame = ttk.LabelFrame(self.root, text="Configurations & File Paths", padding=10)
        top_frame.pack(fill="x", padx=10, pady=6)

        # Row 0: ACT TXT Folder
        ttk.Label(top_frame, text="ACT TXT Output Folder:").grid(row=0, column=0, sticky="w", padx=4, pady=3)
        ttk.Entry(top_frame, textvariable=self.act_folder_path, width=60).grid(row=0, column=1, padx=4, pady=3)
        ttk.Button(top_frame, text="Browse", command=self.browse_act_folder).grid(row=0, column=2, padx=4, pady=3)

        # Row 1: Signal Min/Max File
        ttk.Label(top_frame, text="Signal MinMax File (.xlsx):").grid(row=1, column=0, sticky="w", padx=4, pady=3)
        ttk.Entry(top_frame, textvariable=self.signal_file_path, width=60).grid(row=1, column=1, padx=4, pady=3)
        ttk.Button(top_frame, text="Browse", command=self.browse_signal_file).grid(row=1, column=2, padx=4, pady=3)

        # Calibration Selection (ON/OFF)
        cal_frame = ttk.Frame(top_frame)
        cal_frame.grid(row=1, column=3, padx=12, pady=3, sticky="w")
        ttk.Label(cal_frame, text="Calibration:", font=("Segoe UI", 9, "bold")).pack(side="left", padx=4)
        ttk.Radiobutton(cal_frame, text="OFF (Max of Max)", variable=self.cal_mode, value="OFF", command=self.on_cal_mode_change).pack(side="left", padx=4)
        ttk.Radiobutton(cal_frame, text="ON (Min of Min)", variable=self.cal_mode, value="ON", command=self.on_cal_mode_change).pack(side="left", padx=4)

        # Row 2: Template (.xlsx)
        ttk.Label(top_frame, text="Output Template (.xlsx):").grid(row=2, column=0, sticky="w", padx=4, pady=3)
        ttk.Entry(top_frame, textvariable=self.template_path, width=60).grid(row=2, column=1, padx=4, pady=3)
        ttk.Button(top_frame, text="Browse", command=self.browse_template).grid(row=2, column=2, padx=4, pady=3)

        # Row 3: ALG file (.alg)
        ttk.Label(top_frame, text="Target ALG File (.alg):").grid(row=3, column=0, sticky="w", padx=4, pady=3)
        ttk.Entry(top_frame, textvariable=self.alg_path, width=60).grid(row=3, column=1, padx=4, pady=3)
        ttk.Button(top_frame, text="Browse", command=self.browse_alg).grid(row=3, column=2, padx=4, pady=3)

        # Action: Load Data
        ttk.Button(top_frame, text="Load Channels & ACT Data", command=self.load_all_data).grid(row=0, column=3, rowspan=1, padx=12, pady=3, sticky="ew")

        # Middle Panel: Scrollable Channel Table
        info_lbl = "Rule: C = MAX_OF_MAX if OFF, MIN_OF_MIN if ON. Formulas support A, B, C. Results round up to next integer if value > xx.1"
        table_frame = ttk.LabelFrame(self.root, text=f"Channel Calculations ({info_lbl})", padding=8)
        table_frame.pack(fill="both", expand=True, padx=10, pady=4)

        self.canvas = tk.Canvas(table_frame)
        scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=self.canvas.yview)
        self.scroll_content = ttk.Frame(self.canvas)

        self.scroll_content.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas_window = self.canvas.create_window((0, 0), window=self.scroll_content, anchor="nw")
        self.canvas.configure(yscrollcommand=scrollbar.set)

        self.canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # Bottom Action Bar
        bottom_frame = ttk.Frame(self.root, padding=10)
        bottom_frame.pack(fill="x", padx=12, pady=6)

        ttk.Button(bottom_frame, text="Calculate All Channels", command=self.calculate_all).pack(side="left", padx=6)
        ttk.Button(bottom_frame, text="Process Simple TXT & Fill Template", command=self.fill_simple_thresholds).pack(side="left", padx=6)
        ttk.Button(bottom_frame, text="Export Signals & Finalize Template", command=self.export_to_template).pack(side="left", padx=6)
        ttk.Button(bottom_frame, text="Update .alg File", command=self.update_alg).pack(side="left", padx=6)

        ttk.Label(bottom_frame, textvariable=self.status_msg, font=("Segoe UI", 9, "italic")).pack(side="right", padx=8)

    def browse_act_folder(self):
        d = filedialog.askdirectory()
        if d: self.act_folder_path.set(d)

    def browse_signal_file(self):
        f = filedialog.askopenfilename(filetypes=[("Excel/CSV Files", "*.xlsx *.csv *.xlsm")])
        if f: self.signal_file_path.set(f)

    def browse_template(self):
        f = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")])
        if f: self.template_path.set(f)

    def browse_alg(self):
        f = filedialog.askopenfilename(filetypes=[("ALG Files", "*.alg")])
        if f: self.alg_path.set(f)

    def on_cal_mode_change(self):
        mode = self.cal_mode.get()
        for r in self.channel_rows:
            r.set_mode(mode)

    def load_all_data(self):
        # 1. Process ACT TXT Folder
        folder = self.act_folder_path.get().strip()
        if os.path.exists(folder):
            try:
                self.decided_map, self.raw_val_map, _ = parse_act_txt_folder(folder)
            except Exception as e:
                messagebox.showwarning("ACT Folder Notice", f"Could not parse ACT TXT folder: {e}")

        # 2. Process Signal MinMax File
        sig_file = self.signal_file_path.get().strip()
        if not sig_file or not os.path.exists(sig_file):
            messagebox.showerror("Missing File", "Please select a valid Signal MinMax Excel file.")
            return

        # Clear table
        for w in self.scroll_content.winfo_children():
            w.destroy()
        self.channel_rows.clear()

        try:
            df = pd.read_csv(sig_file) if sig_file.endswith(".csv") else pd.read_excel(sig_file)
            df["MAX"] = pd.to_numeric(df["MAX"], errors="coerce")
            df["MIN"] = pd.to_numeric(df["MIN"], errors="coerce")
            df = df.dropna(subset=["CHANNEL"])

            agg_df = df.groupby("CHANNEL", as_index=False).agg(
                MAX_OF_MAX=("MAX", "max"),
                MIN_OF_MIN=("MIN", "min")
            )

            headers = ["Channel", "MAX_OF_MAX", "MIN_OF_MIN", "Mode", "C", "A", "B", "Formula 1", "Formula 2", "Formula 3", "Result", "Action"]
            for col_idx, h in enumerate(headers):
                ttk.Label(self.scroll_content, text=h, font=("Segoe UI", 8, "bold")).grid(row=0, column=col_idx, padx=3, pady=3)

            current_mode = self.cal_mode.get()
            for idx, row in agg_df.iterrows():
                cr = ChannelRow(self.scroll_content, row["CHANNEL"], row["MAX_OF_MAX"], row["MIN_OF_MIN"], idx + 1, global_mode=current_mode)
                self.channel_rows.append(cr)

            self.status_msg.set(f"Loaded {len(self.channel_rows)} channels successfully.")
        except Exception as e:
            messagebox.showerror("Load Error", f"Failed to load channels:\n{e}")

    def fill_simple_thresholds(self):
        folder = self.act_folder_path.get().strip()
        template = self.template_path.get().strip()
        try:
            if not self.decided_map:
                self.decided_map, self.raw_val_map, _ = parse_act_txt_folder(folder)
            cnt, target = populate_simple_thresholds_template(template, self.decided_map, self.raw_val_map)
            self.status_msg.set(f"Updated {cnt} simple threshold parameters in template.")
            messagebox.showinfo("Success", f"Updated {cnt} simple threshold values in:\n{target}")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to populate simple thresholds:\n{e}")

    def calculate_all(self):
        if not self.channel_rows:
            messagebox.showwarning("No Data", "Please load channels first.")
            return
        cnt = 0
        for r in self.channel_rows:
            res = r.calculate(show_error=False)
            if res["STATUS"] == "OK":
                cnt += 1
        self.status_msg.set(f"Calculated {cnt}/{len(self.channel_rows)} channel formulas (rounded > xx.1).")

    def export_to_template(self):
        if not self.channel_rows:
            messagebox.showwarning("No Data", "Please load channels first.")
            return

        template_file = self.template_path.get().strip()
        if not os.path.exists(template_file):
            messagebox.showerror("Missing Template", f"Template not found:\n{template_file}")
            return

        # Save dialog
        save_file = filedialog.asksaveasfilename(
            title="Save Template With Signal Results",
            initialfile=os.path.basename(template_file),
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")]
        )
        if not save_file:
            return

        # Collect channel calculations
        channel_results = {}
        for r in self.channel_rows:
            res = r.calculate(show_error=False)
            if res["STATUS"] == "OK":
                channel_results[res["CHANNEL"]] = res

        try:
            wb = openpyxl.load_workbook(template_file)
            if "para_selection" not in wb.sheetnames or "parameter_ref" not in wb.sheetnames:
                messagebox.showerror("Invalid Template", "Template must contain 'para_selection' and 'parameter_ref' sheets.")
                return

            ws_sel = wb["para_selection"]
            ws_ref = wb["parameter_ref"]

            # Build mapping from parameter_ref (Canonical Parameter -> tool_para)
            ref_map = {}
            for row in range(2, ws_ref.max_row + 1):
                p_val = ws_ref.cell(row, 3).value   # Col C: Parameter
                tp_val = ws_ref.cell(row, 2).value  # Col B: tool_para
                if p_val and tp_val:
                    clean_p = str(p_val).lstrip("*").strip().upper()
                    ref_map[clean_p] = str(tp_val).strip()

            # Fill para_selection sheet
            updated_count = 0
            for r in range(3, ws_sel.max_row + 1):
                raw_p = ws_sel.cell(r, 2).value  # Col B: Parameter
                if not raw_p:
                    continue

                canon_p = clean_parameter_name(raw_p)
                tool_para = ref_map.get(canon_p.upper())

                if tool_para and tool_para in channel_results:
                    sig = channel_results[tool_para]

                    # Columns F to M
                    ws_sel.cell(r, 6).value = canon_p         # Col F: Parameter
                    ws_sel.cell(r, 7).value = tool_para       # Col G: tool_para
                    ws_sel.cell(r, 8).value = sig.get("C")    # Col H: C
                    ws_sel.cell(r, 9).value = sig.get("A")    # Col I: A
                    ws_sel.cell(r, 10).value = sig.get("B")   # Col J: B
                    ws_sel.cell(r, 11).value = sig.get("FORMULA1") or sig.get("FORMULA", "") # Col K: Formular1
                    if sig.get("FORMULA2"):
                        ws_sel.cell(r, 12).value = sig.get("FORMULA2") # Col L: Formular2
                    if sig.get("FORMULA3"):
                        ws_sel.cell(r, 13).value = sig.get("FORMULA3") # Col M: Formular3

                    # Col N: Compare_simple_value (copies Col D Max Value)
                    simple_val = ws_sel.cell(r, 4).value      # Col D: Max Value
                    ws_sel.cell(r, 14).value = simple_val

                    # Col O: Final_Parameter = max(Simple, Signal_Result)
                    calc_res = sig["RESULT"]
                    if simple_val not in (None, ""):
                        try:
                            ws_sel.cell(r, 15).value = max(float(simple_val), float(calc_res))
                        except Exception:
                            ws_sel.cell(r, 15).value = calc_res
                    else:
                        ws_sel.cell(r, 15).value = calc_res

                    updated_count += 1
                else:
                    # If not calculated via signalminmax, ensure Final_Parameter holds simple thd
                    simple_val = ws_sel.cell(r, 4).value
                    if ws_sel.cell(r, 15).value in (None, "") and simple_val not in (None, ""):
                        ws_sel.cell(r, 15).value = simple_val

            wb.save(save_file)
            self.template_path.set(save_file)
            self.status_msg.set(f"Successfully populated {updated_count} parameters into {os.path.basename(save_file)}")
            messagebox.showinfo("Export Completed", f"Successfully exported {updated_count} parameters into:\n{save_file}\n\nTarget template path updated for ALG generation.")

        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to populate template:\n{e}")

    def update_alg(self):
        template_file = self.template_path.get().strip()
        alg_file = self.alg_path.get().strip()

        if not os.path.exists(template_file):
            messagebox.showerror("Missing Template", f"Template not found:\n{template_file}")
            return
        if not os.path.exists(alg_file):
            messagebox.showerror("Missing ALG File", f"ALG file not found:\n{alg_file}")
            return

        save_file = filedialog.asksaveasfilename(
            title="Save Updated ALG File",
            initialfile=os.path.basename(alg_file).replace(".alg", "_updated.alg"),
            defaultextension=".alg",
            filetypes=[("ALG Files", "*.alg"), ("All Files", "*.*")]
        )
        if not save_file:
            return

        try:
            out_path = update_alg_file(template_file, alg_file, decided_map=self.decided_map, out_alg_file=save_file)
            self.status_msg.set(f"Updated ALG saved to {os.path.basename(out_path)}")
            messagebox.showinfo("Success", f"Successfully updated ALG file:\n{out_path}")
        except Exception as e:
            messagebox.showerror("Update Error", f"Failed to update ALG file:\n{e}")

if __name__ == "__main__":
    root = tk.Tk()
    app = SignalMinMaxTemplateToolApp(root)
    root.mainloop()
