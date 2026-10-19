import os
import re
import math
import pandas as pd
import openpyxl

# -------- Configuration & Constants --------
REQ_OFF = "OFF requirement"
REQ_ON = "ON requirement"

RE_NUMBERS = re.compile(r"\d+")
RE_BASE_INDEX = re.compile(r"(.+?)\[(\d+)\]$")
RE_PREFIX_NUM = re.compile(r"(\d+)_")
RE_PARAM_INDEX = re.compile(r"(.+)\[(\d+)\]$")

RE_SIMPLE_THRESHOLD = re.compile(r"simple\s*threshold\s+([01])", flags=re.IGNORECASE)
RE_CONST_THRESHOLD = re.compile(r"create[_\s]*const[_\s]*threshold\s+([01])", flags=re.IGNORECASE)

RE_ALG_SCALAR = re.compile(r"(\s*)(\w+(\[\d+\])?)\s+([-+]?\d*\.?\d+(?:[eE][-+]?\d+)?)(.*)")
RE_ALG_TABLE_NAME = re.compile(r"NAME\s+(\w+(\[\d+\])?)")
RE_ALG_TABLE_ROW = re.compile(r"^\d+\s+")

def numeric_sort_key(filename):
    numbers = RE_NUMBERS.findall(filename)
    return [int(n) for n in numbers] if numbers else [0]

def extract_base_and_index(filename):
    base = os.path.splitext(filename)[0]
    parts = base.split("_", 1)
    name_with_idx = parts[1] if len(parts) == 2 else base
    m = RE_BASE_INDEX.match(name_with_idx)
    return (m.group(1), int(m.group(2))) if m else (name_with_idx, None)

def round_threshold_value(val):
    """Round up to the next integer if fractional value is > 0.1."""
    if pd.isna(val) or val is None:
        return val
    try:
        f = float(val)
        floor_val = math.floor(f)
        if (f - floor_val) > 0.1000000001:
            return float(math.ceil(f))
        return float(floor_val)
    except Exception:
        return val

def extract_parameter_name(filename):
    base = os.path.splitext(filename)[0]
    return re.sub(r"^\d+_", "", base).strip()

def extract_prefix_number(filename):
    base = os.path.splitext(filename)[0]
    m = RE_PREFIX_NUM.match(base)
    return int(m.group(1)) if m else None

def get_requirement_from_lines(lines):
    const_val, simple_val = None, None
    for line in lines:
        mc = RE_CONST_THRESHOLD.search(line)
        if mc: const_val = int(mc.group(1))
        ms = RE_SIMPLE_THRESHOLD.search(line)
        if ms: simple_val = int(ms.group(1))
    return REQ_ON if const_val == 1 else REQ_OFF

def extract_txt_value(lines):
    for idx in [19, 20]:
        if len(lines) > idx:
            parts = lines[idx].strip().split()
            if parts and parts[0] == "0" and len(parts) > 1:
                try:
                    return float(parts[1])
                except Exception:
                    pass
    return None

def format_val_for_alg(val):
    if pd.isna(val) or val == "":
        return ""
    try:
        f = float(val)
        return str(int(f)) if f.is_integer() else str(f)
    except Exception:
        return str(val)

def normalize_param_name(raw):
    if not raw or pd.isna(raw):
        return ""
    s = re.sub(r"^\d+_", "", str(raw)).strip()
    return re.sub(r"_(OFF|ON_LEFT|ON_RIGHT|ON)$", "", s, flags=re.IGNORECASE).strip()

def parse_act_txt_folder(folder):
    """Parses TXT output from ACT folder and applies pair decision logic."""
    if not os.path.exists(folder):
        raise FileNotFoundError(f"Folder not found: {folder}")

    txt_files = sorted([f for f in os.listdir(folder) if f.lower().endswith(".txt")], key=numeric_sort_key)
    if not txt_files:
        raise ValueError(f"No .txt files found in: {folder}")

    file_info = []
    file_requirements = {}
    for file in txt_files:
        try:
            with open(os.path.join(folder, file), "r") as fp:
                lines = fp.readlines()
                val = extract_txt_value(lines)
                req = get_requirement_from_lines(lines)
        except Exception:
            val, req = None, REQ_OFF

        base, idx = extract_base_and_index(file)
        param = extract_parameter_name(file)
        file_info.append({"Filename": file, "Base": base, "Index": idx, "Value": val, "Param": param})
        file_requirements[file] = req

    grouped = {}
    decided_map = {}
    raw_val_map = {}

    for e in file_info:
        decided_map[e["Filename"]] = e["Value"]
        raw_val_map[e["Filename"]] = e["Value"]
        base, idx = e["Base"], e["Index"]
        grouped.setdefault(base, {})
        if idx is not None:
            grouped[base][idx] = e

    # Pair comparison [0]&[1], [2]&[3], etc.
    for base, idx_dict in grouped.items():
        for i in [0, 2, 4]:
            e1, e2 = idx_dict.get(i), idx_dict.get(i + 1)
            if e1 and e2 and e1["Value"] is not None and e2["Value"] is not None:
                req1 = file_requirements.get(e1["Filename"], REQ_OFF)
                req2 = file_requirements.get(e2["Filename"], REQ_OFF)
                req = REQ_OFF if req1 != req2 else req1
                chosen = max(e1["Value"], e2["Value"]) if req == REQ_OFF else min(e1["Value"], e2["Value"])
                decided_map[e1["Filename"]] = chosen
                decided_map[e2["Filename"]] = chosen

    # Suffix pairs (1L/1R, 2L/2R, FL/FR, ML/MR, RL/RR)
    suffix_pairs = [("1L", "1R"), ("2L", "2R"), ("FL", "FR"), ("ML", "MR"), ("RL", "RR")]
    base_to_entry = {str(e["Base"]).upper(): e for e in file_info}
    for left, right in suffix_pairs:
        for e in file_info:
            base = e["Base"]
            if base.upper().endswith(left):
                base_right = base[:-len(left)] + right
                e_right = base_to_entry.get(base_right.upper())
                if e_right and e["Value"] is not None and e_right["Value"] is not None:
                    chosen = max(e["Value"], e_right["Value"])
                    decided_map[e["Filename"]] = chosen
                    decided_map[e_right["Filename"]] = chosen

    return decided_map, raw_val_map, file_info

def populate_simple_thresholds_template(template_file, decided_map, raw_val_map, output_template=None):
    """Fills Cols C, D, N, and default O in sheet 'para_selection'."""
    if not os.path.exists(template_file):
        raise FileNotFoundError(f"Template not found: {template_file}")

    wb = openpyxl.load_workbook(template_file)
    if "para_selection" not in wb.sheetnames:
        raise ValueError("Sheet 'para_selection' not found in template.")

    ws = wb["para_selection"]
    matched_count = 0
    fn_lookup = {os.path.splitext(f)[0].upper(): f for f in decided_map.keys()}

    for r in range(3, ws.max_row + 1):
        raw_param = ws.cell(r, 2).value  # Col B
        if not raw_param:
            continue

        raw_str = str(raw_param).strip()
        fn = fn_lookup.get(raw_str.upper()) or fn_lookup.get(raw_str.upper() + ".TXT")

        if fn:
            raw_val = raw_val_map.get(fn)
            dec_val = decided_map.get(fn)
            ws.cell(r, 3).value = raw_val    # Col C: Value
            ws.cell(r, 4).value = dec_val    # Col D: Max Value
            ws.cell(r, 14).value = dec_val   # Col N: Compare_simple_value

            current_final = ws.cell(r, 15).value
            if current_final in (None, ""):
                ws.cell(r, 15).value = dec_val
            matched_count += 1

    target_file = output_template or template_file
    wb.save(target_file)
    return matched_count, target_file

def update_alg_file(template_file, alg_file, decided_map=None, out_alg_file=None):
    """Reads final parameters from the template and replaces values in target .alg."""
    if not os.path.exists(template_file):
        raise FileNotFoundError(f"Template not found: {template_file}")
    if not os.path.exists(alg_file):
        raise FileNotFoundError(f"ALG file not found: {alg_file}")

    wb = openpyxl.load_workbook(template_file, data_only=True)
    if "para_selection" not in wb.sheetnames:
        raise ValueError("Sheet 'para_selection' missing from template.")
    ws_sel = wb["para_selection"]

    numeric_param_map = {}
    param_map = {}

    # 1. Read Final_Parameter (Col O) first, then fallback to Max Value (Col D)
    for r in range(3, ws_sel.max_row + 1):
        raw_p = ws_sel.cell(r, 2).value
        final_val = ws_sel.cell(r, 15).value  # Col O: Final_Parameter
        max_val = ws_sel.cell(r, 4).value    # Col D: Max Value (Simple Thd)

        val_to_use = final_val if final_val not in (None, "") else max_val
        if raw_p and val_to_use not in (None, ""):
            canon_p = normalize_param_name(str(raw_p))
            try:
                num = float(val_to_use)
                numeric_param_map[canon_p] = max(numeric_param_map.get(canon_p, num), num)
            except (ValueError, TypeError):
                if canon_p not in param_map:
                    param_map[canon_p] = str(val_to_use)

    for k, v in numeric_param_map.items():
        param_map[k] = format_val_for_alg(v)

    # 2. Add fallback simple threshold parameters from decided_map if provided
    if decided_map:
        for fn, v in decided_map.items():
            p = extract_parameter_name(fn)
            p_norm = normalize_param_name(p)
            if p_norm not in param_map and v is not None:
                param_map[p_norm] = format_val_for_alg(v)

    # 3. Forced copy: 79_SacEcuYThdLow2L -> SacEcuYThdLow1L, 81 -> 82
    if "SacEcuYThdLow2L" in param_map:
        param_map["SacEcuYThdLow1L"] = param_map["SacEcuYThdLow2L"]
    if "SacEcuYThdLow2R" in param_map:
        param_map["SacEcuYThdLow1R"] = param_map["SacEcuYThdLow2R"]

    # 4. Explicit synthetic cases
    low_vals = [float(param_map[k]) for k in ["SacEcuYThdLow2L", "SacEcuYThdLow1L", "SacEcuYThdLow2R", "SacEcuYThdLow1R"] if k in param_map]
    if low_vals:
        param_map["SprPlsEcuYLpShortThd"] = format_val_for_alg(max(low_vals) + 1)

    # 5. Inject into .alg
    with open(alg_file, "r") as fp:
        lines = fp.readlines()

    param_map_lower = {k.lower(): v for k, v in param_map.items()}

    new_lines, i = [], 0
    while i < len(lines):
        line = lines[i]
        m_scalar = RE_ALG_SCALAR.match(line)
        if m_scalar:
            p_name = m_scalar.group(2)
            matched_val = param_map.get(p_name) or param_map_lower.get(p_name.lower())
            if matched_val is not None:
                new_lines.append(f"{m_scalar.group(1)}{p_name} {matched_val}{m_scalar.group(5)}\n")
            else:
                new_lines.append(line)
            i += 1
            continue

        m_tbl = RE_ALG_TABLE_NAME.match(line)
        if m_tbl:
            p_name = m_tbl.group(1)
            new_lines.append(line)
            i += 1
            val = param_map.get(p_name) or param_map_lower.get(p_name.lower())
            while i < len(lines) and not RE_ALG_TABLE_ROW.match(lines[i]):
                new_lines.append(lines[i])
                i += 1
            while i < len(lines) and RE_ALG_TABLE_ROW.match(lines[i]):
                idx = lines[i].split()[0]
                new_lines.append(f"{idx}\t{val}\n" if val is not None else lines[i])
                i += 1
            continue

        new_lines.append(line)
        i += 1

    output_path = out_alg_file or alg_file.replace(".alg", "_updated.alg")
    with open(output_path, "w") as fp:
        fp.writelines(new_lines)

    return output_path